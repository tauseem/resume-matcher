from pypdf import PdfReader
from docx import Document


def extract_text(uploaded_file):
    name = uploaded_file.name.lower()

    if name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n".join(pages)

    if name.endswith(".docx"):
        document = Document(uploaded_file)
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    lines.append(cell.text)
        return "\n".join(lines)

    raise ValueError("Only PDF and DOCX files are supported.")

